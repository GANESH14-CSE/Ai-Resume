import io
import re

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import pypdf
except ImportError:
    pypdf = None

from docx import Document

def extract_text_from_file(file_obj, filename: str = '') -> str:
    """
    Extracts plain text from an uploaded file (PDF, DOCX, TXT).
    """
    filename_lower = filename.lower()

    # Read bytes
    content = file_obj.read() if hasattr(file_obj, 'read') else file_obj
    if isinstance(content, str):
        return content

    # 1. PDF via PyMuPDF or pypdf
    if filename_lower.endswith('.pdf') or content.startswith(b'%PDF'):
        if fitz:
            try:
                doc = fitz.open(stream=content, filetype="pdf")
                pages_text = []
                for page in doc:
                    pages_text.append(page.get_text())
                full_text = "\n".join(pages_text)
                return re.sub(r'\s+', ' ', full_text).strip()
            except Exception:
                pass
        if pypdf:
            try:
                reader = pypdf.PdfReader(io.BytesIO(content))
                pages_text = [page.extract_text() or '' for page in reader.pages]
                full_text = "\n".join(pages_text)
                return re.sub(r'\s+', ' ', full_text).strip()
            except Exception:
                pass
        return content.decode('utf-8', errors='ignore')

    # 2. DOCX via python-docx
    if filename_lower.endswith('.docx'):
        try:
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
        except Exception:
            return content.decode('utf-8', errors='ignore')

    # 3. Plain text / fallback
    try:
        return content.decode('utf-8').strip()
    except UnicodeDecodeError:
        return content.decode('latin-1', errors='ignore').strip()
